"""
Twitter / X Profile Finder — single-file deployment version
All modules inlined so Streamlit Cloud needs no src/ folder.
"""

# ── stdlib + third-party imports (all must be in requirements.txt) ────────────
import io
import re
import sys
import os
import time
import json
import logging
from dataclasses import dataclass, field
from typing import Optional, Union, IO

import streamlit as st
import pandas as pd
import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# =============================================================================
# DATA MODELS
# =============================================================================

@dataclass
class Person:
    first_name: str
    last_name: str
    address: str = ""
    employer: str = ""

    @property
    def full_name(self) -> str:
        return f"{self.first_name.strip()} {self.last_name.strip()}".strip()

    @property
    def city(self) -> str:
        parts = [p.strip() for p in self.address.split(",")]
        if len(parts) >= 2:
            return parts[-2]
        return parts[0] if parts else ""

    @property
    def state(self) -> str:
        US_STATES = {
            "AL","AK","AZ","AR","CA","CO","CT","DE","FL","GA","HI","ID","IL","IN",
            "IA","KS","KY","LA","ME","MD","MA","MI","MN","MS","MO","MT","NE","NV",
            "NH","NJ","NM","NY","NC","ND","OH","OK","OR","PA","RI","SC","SD","TN",
            "TX","UT","VT","VA","WA","WV","WI","WY","DC",
        }
        parts = [p.strip() for p in self.address.split(",")]
        last = parts[-1] if parts else ""
        for t in last.split():
            if t.upper() in US_STATES:
                return t.upper()
        return ""


@dataclass
class TwitterProfile:
    handle: str
    display_name: str
    bio: str
    location: str
    followers: int
    verified: bool
    created_year: int
    profile_url: str
    avatar_url: str = ""
    source: str = "search"


@dataclass
class MatchResult:
    person: Person
    profile: Optional[TwitterProfile]
    score: int
    confidence: str
    score_breakdown: dict = field(default_factory=dict)
    search_query: str = ""
    error: str = ""
    keyword_hits: dict = field(default_factory=dict)   # keyword -> ["bio", "tweet: ..."]
    keyword_flagged: bool = False

    @property
    def emoji(self) -> str:
        return {"HIGH": "✅", "MEDIUM": "🟡", "LOW": "🟠", "NO MATCH": "❌"}.get(self.confidence, "❓")


# =============================================================================
# CONFIDENCE SCORER
# =============================================================================

def _normalise(text: str) -> str:
    return re.sub(r"[^a-z0-9 ]", " ", text.lower())


def score_match(person: Person, profile: TwitterProfile):
    breakdown = {}
    total = 0

    full       = person.full_name.lower()
    disp       = profile.display_name.lower()
    handle_c   = profile.handle.lstrip("@").lower()
    bio_norm   = _normalise(profile.bio)
    loc_norm   = _normalise(profile.location)
    combined   = bio_norm + " " + loc_norm
    first_l    = person.first_name.lower()
    last_l     = person.last_name.lower()

    # Name
    if full in disp or disp in full:
        breakdown["name_exact"] = 40;      total += 40
    elif last_l in disp and first_l in disp:
        breakdown["name_both_parts"] = 35; total += 35
    elif last_l in disp or last_l in handle_c:
        breakdown["name_last_only"] = 20;  total += 20
    elif first_l in disp:
        breakdown["name_first_only"] = 10; total += 10

    # Employer
    if person.employer:
        emp_words = [w for w in _normalise(person.employer).split() if len(w) > 3]
        matches = sum(1 for w in emp_words if w in bio_norm)
        if emp_words and matches / len(emp_words) >= 0.5:
            breakdown["employer_match"] = 25;   total += 25
        elif matches > 0:
            breakdown["employer_partial"] = 10; total += 10

    # City
    if person.city:
        city_norm = _normalise(person.city)
        if city_norm and city_norm in combined:
            breakdown["city_match"] = 20; total += 20

    # State
    if person.state and person.state.lower() in combined:
        breakdown["state_match"] = 10; total += 10

    # Bio keywords
    keywords = set()
    for src in [person.employer, person.city]:
        for word in _normalise(src).split():
            if len(word) > 4:
                keywords.add(word)
    kw_hits = sum(1 for kw in keywords if kw in bio_norm)
    if kw_hits:
        pts = min(kw_hits * 5, 15)
        breakdown["bio_keywords"] = pts; total += pts

    # Verified / age
    if profile.verified:
        breakdown["verified"] = 5; total += 5
    if profile.created_year and profile.created_year <= 2022:
        breakdown["account_age"] = 5; total += 5

    total = min(total, 100)
    if total >= 75:   confidence = "HIGH"
    elif total >= 45: confidence = "MEDIUM"
    elif total >= 20: confidence = "LOW"
    else:             confidence = "NO MATCH"

    return total, breakdown


# =============================================================================
# SEARCH
# =============================================================================

def _build_queries(person: Person):
    queries = []
    name = person.full_name
    city = person.city
    emp  = person.employer
    if emp and city:
        queries.append(f'"{name}" "{emp}" "{city}" site:twitter.com OR site:x.com')
        queries.append(f'"{name}" "{emp}" site:twitter.com OR site:x.com')
    if emp:
        queries.append(f'"{name}" "{emp}" twitter')
    if city:
        queries.append(f'"{name}" "{city}" twitter')
    queries.append(f'"{name}" twitter')
    queries.append(f'{name} site:twitter.com')
    return queries


def _parse_twitter_url(url: str):
    m = re.search(r'(?:twitter\.com|x\.com)/([A-Za-z0-9_]{1,50})(?:/|$|\?)', url)
    if m:
        handle = m.group(1)
        if handle.lower() not in ("intent","search","hashtag","share","home",
                                   "explore","notifications","messages","i"):
            return "@" + handle
    return None


def _build_profile(candidate: dict) -> TwitterProfile:
    raw    = candidate.get("raw", {})
    handle = candidate.get("handle", "")
    if raw:
        return TwitterProfile(
            handle=handle,
            display_name=raw.get("display_name","") or candidate.get("title",""),
            bio=raw.get("bio","") or candidate.get("snippet",""),
            location=raw.get("location",""),
            followers=int(raw.get("followers", 0) or 0),
            verified=bool(raw.get("verified", False)),
            created_year=int(raw.get("created_year", 0) or 0),
            profile_url=raw.get("profile_url","") or candidate.get("url",""),
            source="claude_search",
        )
    title = candidate.get("title","")
    display_name = re.sub(r'\s*\(@[^)]+\)', '', title).split(" | ")[0].strip()
    return TwitterProfile(
        handle=handle,
        display_name=display_name,
        bio=candidate.get("snippet",""),
        location="",
        followers=0,
        verified=False,
        created_year=0,
        profile_url=candidate.get("url", f"https://twitter.com/{handle.lstrip('@')}"),
        source="search_snippet",
    )


def _search_claude(person: Person, anthropic_key: str):
    import anthropic as _anthropic
    client = _anthropic.Anthropic(api_key=anthropic_key)
    prompt = f"""Find the Twitter/X profile for this person using web search.

Person details:
- Full name: {person.full_name}
- Address/City: {person.address}
- Employer: {person.employer}

Search for their Twitter or X.com profile. Try multiple searches if needed.
Return ONLY a JSON array of candidate profiles. Each object must have:
  handle, display_name, bio, location, followers (int), verified (bool),
  created_year (int), profile_url, confidence_notes

Return ONLY the JSON array, no markdown, no other text. If nothing found return [].
Try at least 3 different search queries before giving up."""

    candidates = []
    seen = set()
    try:
        resp = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2000,
            tools=[{"type": "web_search_20250305", "name": "web_search"}],
            messages=[{"role": "user", "content": prompt}],
        )
        full_text = "".join(b.text for b in resp.content if hasattr(b, "text"))
        clean = re.sub(r"```(?:json)?", "", full_text).strip().strip("`").strip()
        m = re.search(r'\[.*\]', clean, re.DOTALL)
        if m:
            for item in json.loads(m.group(0)):
                handle = item.get("handle","").strip()
                if not handle.startswith("@"):
                    handle = "@" + handle
                if handle and handle not in seen and handle != "@":
                    seen.add(handle)
                    candidates.append({
                        "handle": handle,
                        "title": item.get("display_name",""),
                        "snippet": item.get("bio",""),
                        "url": item.get("profile_url", f"https://twitter.com/{handle.lstrip('@')}"),
                        "query": "claude_web_search",
                        "raw": item,
                    })
    except Exception as e:
        logger.warning(f"Claude search error for {person.full_name}: {e}")
    return candidates


def _search_serpapi(person: Person, api_key: str):
    candidates = []
    seen = set()
    for query in _build_queries(person)[:3]:
        try:
            resp = requests.get(
                "https://serpapi.com/search",
                params={"q": query, "api_key": api_key, "num": 5, "engine": "google"},
                timeout=10,
            )
            if resp.status_code != 200:
                continue
            for r in resp.json().get("organic_results", []):
                url = r.get("link","")
                handle = _parse_twitter_url(url)
                if handle and handle not in seen:
                    seen.add(handle)
                    candidates.append({"handle": handle, "title": r.get("title",""),
                                       "snippet": r.get("snippet",""), "url": url, "query": query})
            time.sleep(0.3)
        except Exception as e:
            logger.warning(f"SerpAPI: {e}")
    return candidates


def find_twitter(person: Person, anthropic_key="", serpapi_key="") -> MatchResult:
    candidates = []
    if anthropic_key:
        candidates.extend(_search_claude(person, anthropic_key))
    if not candidates and serpapi_key:
        candidates.extend(_search_serpapi(person, serpapi_key))
    if not candidates:
        return MatchResult(person=person, profile=None, score=0, confidence="NO MATCH",
                           error="No results returned.")

    scored = []
    for cand in candidates:
        profile = _build_profile(cand)
        score, breakdown = score_match(person, profile)
        scored.append((score, breakdown, profile, cand.get("query","")))
    scored.sort(key=lambda x: x[0], reverse=True)
    best_score, best_breakdown, best_profile, best_query = scored[0]

    if best_score >= 75:   confidence = "HIGH"
    elif best_score >= 45: confidence = "MEDIUM"
    elif best_score >= 20: confidence = "LOW"
    else:                  confidence = "NO MATCH"

    return MatchResult(person=person, profile=best_profile, score=best_score,
                       confidence=confidence, score_breakdown=best_breakdown,
                       search_query=best_query)


# =============================================================================
# KEYWORD SCANNER
# =============================================================================

def scan_keywords(result: MatchResult, keywords: list, anthropic_key: str) -> MatchResult:
    """
    Given a MatchResult with a confirmed profile, ask Claude to fetch that
    person's recent tweets and check both bio + tweets for each keyword.
    Populates result.keyword_hits and result.keyword_flagged in-place.
    """
    if not result.profile or not keywords or not anthropic_key:
        return result

    import anthropic as _anthropic
    client = _anthropic.Anthropic(api_key=anthropic_key)

    handle   = result.profile.handle
    bio_text = result.profile.bio or ""
    kw_list  = ", ".join(f'"{k}"' for k in keywords)

    prompt = f"""You are scanning a Twitter/X profile for specific keywords.

Profile: {handle}
Bio text already retrieved: "{bio_text}"

Your tasks:
1. Search the web for recent tweets by {handle} — try queries like:
   - site:twitter.com {handle} tweets
   - {handle} twitter recent posts
   Collect up to 20 recent tweet texts if available.

2. Check BOTH the bio AND the tweets for these keywords (case-insensitive): {kw_list}

3. Return ONLY a JSON object with this exact structure:
{{
  "bio_hits": ["keyword1", "keyword2"],
  "tweet_hits": {{
    "keyword1": ["exact tweet snippet containing keyword1 (max 100 chars)"],
    "keyword2": ["exact tweet snippet..."]
  }},
  "tweets_found": true
}}

- bio_hits: list of keywords found anywhere in the bio
- tweet_hits: for each keyword found in tweets, list up to 3 short snippets as evidence
- tweets_found: true if you found any recent tweets, false if private or inactive

Return ONLY the JSON, no markdown, no other text."""

    try:
        resp = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=1500,
            tools=[{"type": "web_search_20250305", "name": "web_search"}],
            messages=[{"role": "user", "content": prompt}],
        )
        full_text = "".join(b.text for b in resp.content if hasattr(b, "text"))
        clean = re.sub(r"```(?:json)?", "", full_text).strip().strip("`").strip()
        m = re.search(r'\{.*\}', clean, re.DOTALL)
        if m:
            data = json.loads(m.group(0))
            hits = {}
            for kw in data.get("bio_hits", []):
                kw_l = kw.lower()
                hits.setdefault(kw_l, [])
                if "bio" not in hits[kw_l]:
                    hits[kw_l].append("bio")
            for kw, snippets in data.get("tweet_hits", {}).items():
                kw_l = kw.lower()
                hits.setdefault(kw_l, [])
                for s in snippets[:3]:
                    hits[kw_l].append(f'tweet: "{str(s)[:100]}"')
            result.keyword_hits    = hits
            result.keyword_flagged = len(hits) > 0
    except Exception as e:
        logger.warning(f"Keyword scan error for {handle}: {e}")

    return result


# =============================================================================
# SPREADSHEET I/O
# =============================================================================

FILLS = {
    "HIGH":     PatternFill("solid", start_color="DCFCE7"),
    "MEDIUM":   PatternFill("solid", start_color="FEF9C3"),
    "LOW":      PatternFill("solid", start_color="FFEDD5"),
    "NO MATCH": PatternFill("solid", start_color="FEE2E2"),
    "HEADER":   PatternFill("solid", start_color="1E293B"),
}
CONFIDENCE_LABELS = {
    "HIGH": "✓ HIGH", "MEDIUM": "~ MEDIUM", "LOW": "! LOW", "NO MATCH": "✗ NO MATCH"
}
_thin   = Side(style="thin", color="E2E8F0")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _find_col(columns, keywords):
    for col in columns:
        if any(kw in col.lower() for kw in keywords):
            return col
    return None


def load_persons(file):
    warnings = []
    if hasattr(file, "name") and file.name.endswith(".csv"):
        df = pd.read_csv(file)
    else:
        df = pd.read_excel(file)
    cols = list(df.columns)
    first_col    = _find_col(cols, ["first"])
    last_col     = _find_col(cols, ["last"])
    address_col  = _find_col(cols, ["address","addr","street"])
    employer_col = _find_col(cols, ["employer","company","organization","org","work"])
    if not first_col:
        raise ValueError("Could not find a 'First Name' column.")
    if not last_col:
        raise ValueError("Could not find a 'Last Name' column.")
    if not address_col:
        warnings.append("No address column found — city/state matching disabled.")
    if not employer_col:
        warnings.append("No employer column found — employer matching disabled.")
    persons = []
    for _, row in df.iterrows():
        persons.append(Person(
            first_name=str(row.get(first_col,"") or "").strip(),
            last_name=str(row.get(last_col,"") or "").strip(),
            address=str(row.get(address_col,"") or "").strip() if address_col else "",
            employer=str(row.get(employer_col,"") or "").strip() if employer_col else "",
        ))
    return df, persons, warnings


def write_results(original_file, results):
    if hasattr(original_file, "seek"):
        original_file.seek(0)
    try:
        if hasattr(original_file, "name") and original_file.name.endswith(".csv"):
            df = pd.read_csv(original_file)
            wb = Workbook(); ws = wb.active; ws.title = "Results"
            ws.append(list(df.columns))
            for row in df.itertuples(index=False): ws.append(list(row))
        else:
            wb = load_workbook(original_file); ws = wb.active
    except Exception:
        wb = Workbook(); ws = wb.active; ws.title = "Results"

    orig_cols = ws.max_column
    new_headers = ["Twitter Handle","Profile URL","Confidence","Score (0-100)",
                   "Display Name","Bio","Location","Score Breakdown",
                   "Keyword Flagged","Keywords Matched","Keyword Evidence"]
    hfont = Font(color="FFFFFF", bold=True, name="Calibri", size=10)

    for c in range(1, orig_cols + 1):
        cell = ws.cell(1, c)
        cell.fill = FILLS["HEADER"]; cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for i, h in enumerate(new_headers, start=orig_cols + 1):
        cell = ws.cell(1, i, h)
        cell.fill = FILLS["HEADER"]; cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[1].height = 28

    for row_idx, result in enumerate(results, start=2):
        fill  = FILLS.get(result.confidence, FILLS["NO MATCH"])
        dfont = Font(name="Calibri", size=10)
        for c in range(1, orig_cols + 1):
            cell = ws.cell(row_idx, c)
            cell.fill = fill; cell.font = dfont; cell.border = _BORDER
            cell.alignment = Alignment(vertical="center")
        p = result.profile
        breakdown_str = " | ".join(f"{k}: +{v}" for k, v in result.score_breakdown.items())
        kw_flagged  = "YES" if result.keyword_flagged else ("N/A" if not result.keyword_hits and not result.profile else "NO")
        kw_matched  = ", ".join(result.keyword_hits.keys()) if result.keyword_hits else ""
        kw_evidence = " | ".join(
            f"{kw}: {'; '.join(locs)}" for kw, locs in result.keyword_hits.items()
        ) if result.keyword_hits else ""
        vals = [
            p.handle if p else "",
            p.profile_url if p else "",
            CONFIDENCE_LABELS.get(result.confidence, result.confidence),
            result.score,
            p.display_name if p else "",
            (p.bio[:200] + "…") if p and len(p.bio) > 200 else (p.bio if p else ""),
            p.location if p else "",
            breakdown_str,
            kw_flagged,
            kw_matched,
            kw_evidence,
        ]
        for i, val in enumerate(vals, start=orig_cols + 1):
            cell = ws.cell(row_idx, i, val)
            cell.fill = fill; cell.font = dfont; cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(i > orig_cols + 2))
        ws.row_dimensions[row_idx].height = 18

    for offset, width in enumerate([18,35,14,12,20,45,20,55,10,30,60], start=1):
        ws.column_dimensions[get_column_letter(orig_cols + offset)].width = width

    if "Summary" in wb.sheetnames: del wb["Summary"]
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 12
    ws2.append(["Confidence Level","Count","% of Total"])
    for c in range(1, 4):
        cell = ws2.cell(1, c)
        cell.fill = FILLS["HEADER"]; cell.font = hfont
        cell.alignment = Alignment(horizontal="center"); cell.border = _BORDER
    counts = {"HIGH":0,"MEDIUM":0,"LOW":0,"NO MATCH":0}
    for r in results: counts[r.confidence] = counts.get(r.confidence,0) + 1
    total = len(results)
    for level in ["HIGH","MEDIUM","LOW","NO MATCH"]:
        n = counts[level]
        ws2.append([CONFIDENCE_LABELS[level], n, f"{100*n/total:.1f}%" if total else "0%"])
        for c in range(1, 4):
            cell = ws2.cell(ws2.max_row, c)
            cell.fill = FILLS[level]; cell.font = Font(name="Calibri", size=10)
            cell.border = _BORDER; cell.alignment = Alignment(horizontal="center")

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# =============================================================================
# STREAMLIT UI
# =============================================================================

st.set_page_config(page_title="Twitter Profile Finder", page_icon="🔍", layout="wide",
                   initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:wght@300;400;500;600&display=swap');
html, body, [class*="css"] { font-family: 'IBM Plex Sans', sans-serif; }
.stApp { background: #0a0f1e; color: #e2e8f0; }
.hero-wrap {
    border: 1px solid #1e293b; border-radius: 12px;
    padding: 2.5rem 2rem 2rem;
    background: linear-gradient(135deg, #0f172a 0%, #0a0f1e 100%);
    margin-bottom: 1.5rem;
}
.hero-title { font-family: 'IBM Plex Mono', monospace; font-size: 1.7rem; font-weight: 500; color: #f1f5f9; letter-spacing: -0.5px; margin: 0 0 0.4rem; }
.hero-sub { color: #64748b; font-size: 0.9rem; margin: 0; font-weight: 300; }
.tag { display: inline-block; background: #0ea5e920; color: #38bdf8; border: 1px solid #0ea5e940; border-radius: 4px; padding: 1px 8px; font-size: 0.72rem; font-family: 'IBM Plex Mono', monospace; margin-bottom: 1rem; }
.metric-row { display: flex; gap: 0.75rem; flex-wrap: wrap; margin: 1rem 0; }
.metric { flex: 1; min-width: 110px; background: #0f172a; border: 1px solid #1e293b; border-radius: 8px; padding: 1rem 0.75rem; text-align: center; }
.metric .num { font-family: 'IBM Plex Mono', monospace; font-size: 1.8rem; font-weight: 500; line-height: 1; margin-bottom: 4px; }
.metric .lbl { font-size: 0.7rem; text-transform: uppercase; letter-spacing: 1px; color: #475569; }
.result-row { display: flex; align-items: center; gap: 0.75rem; padding: 0.6rem 1rem; border-radius: 6px; border: 1px solid #1e293b; margin: 0.3rem 0; background: #0f172a; }
.badge { font-family: 'IBM Plex Mono', monospace; font-size: 0.65rem; padding: 2px 7px; border-radius: 3px; font-weight: 500; white-space: nowrap; min-width: 72px; text-align: center; }
.badge-HIGH     { background: #14532d40; color: #4ade80; border: 1px solid #16a34a50; }
.badge-MEDIUM   { background: #71350040; color: #fbbf24; border: 1px solid #d9770050; }
.badge-LOW      { background: #7c2d1240; color: #fb923c; border: 1px solid #ea580c50; }
.badge-NOMATCH  { background: #7f1d1d40; color: #f87171; border: 1px solid #dc262650; }
.score-bar-wrap { flex: 1; background: #1e293b; border-radius: 3px; height: 4px; min-width: 60px; }
.score-bar { height: 4px; border-radius: 3px; }
.person-name { font-weight: 500; font-size: 0.9rem; min-width: 160px; color: #e2e8f0; }
.handle-link { font-family: 'IBM Plex Mono', monospace; font-size: 0.8rem; color: #38bdf8; text-decoration: none; }
.score-num { font-family: 'IBM Plex Mono', monospace; font-size: 0.75rem; color: #475569; min-width: 28px; text-align: right; }
.info-box { background: #0c1a2e; border-left: 3px solid #0ea5e9; border-radius: 0 6px 6px 0; padding: 0.6rem 0.9rem; font-size: 0.82rem; color: #64748b; margin: 0.5rem 0; }
.warn-box { background: #1a1200; border-left: 3px solid #d97706; border-radius: 0 6px 6px 0; padding: 0.6rem 0.9rem; font-size: 0.82rem; color: #92400e; margin: 0.4rem 0; }
.breakdown-pill { display: inline-block; background: #1e293b; color: #94a3b8; border-radius: 4px; padding: 1px 6px; font-family: 'IBM Plex Mono', monospace; font-size: 0.68rem; margin: 1px; }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero-wrap">
    <div class="tag">v1.1 · profile finder + keyword scanner</div>
    <div class="hero-title">Twitter / X Profile Finder</div>
    <p class="hero-sub">Upload a spreadsheet of names, addresses, and employers —
    get back matched Twitter profiles with confidence scores, then scan bios and tweets for keywords.</p>
</div>
""", unsafe_allow_html=True)

# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### ⚙️ Configuration")
    st.markdown("---")

    # Load from Streamlit secrets if available, otherwise show input
    default_key = ""
    try:
        default_key = st.secrets.get("ANTHROPIC_API_KEY", "")
    except Exception:
        pass

    if default_key:
        anthropic_key = default_key
        st.success("✓ Anthropic API key loaded from secrets")
    else:
        anthropic_key = st.text_input("Anthropic API Key", type="password",
                                       placeholder="sk-ant-...",
                                       help="Get one at console.anthropic.com")

    serpapi_key = st.text_input("SerpAPI Key (optional)", type="password",
                                 placeholder="Fallback search backend")
    st.markdown("---")
    st.markdown("""#### Confidence levels
<div style='font-size:0.8rem;color:#475569;line-height:2'>
<span style='color:#4ade80'>✓ HIGH ≥ 75</span> — Strong match<br>
<span style='color:#fbbf24'>~ MED 45–74</span> — Likely match<br>
<span style='color:#fb923c'>! LOW 20–44</span> — Possible match<br>
<span style='color:#f87171'>✗ NONE &lt;20</span> — Not found
</div>""", unsafe_allow_html=True)
    st.markdown("---")
    st.markdown("""<div style='font-size:0.75rem;color:#334155'>
<b>Scoring signals</b><br>
+40 Exact name match<br>
+25 Employer in bio<br>
+20 City in profile<br>
+10 State in profile<br>
+15 Bio keyword overlap<br>
+5 Verified account<br>
+5 Account age > 2yr
</div>""", unsafe_allow_html=True)

# ── Upload ────────────────────────────────────────────────────────────────────
st.markdown("### 1. Upload your spreadsheet")
st.markdown('<div class="info-box">Required columns: <b>First Name</b>, <b>Last Name</b> &nbsp;·&nbsp; Recommended: <b>Address</b>, <b>Employer</b></div>', unsafe_allow_html=True)
uploaded = st.file_uploader("", type=["xlsx","csv"], label_visibility="collapsed")

if uploaded:
    try:
        df_orig, persons, warnings = load_persons(uploaded)
    except ValueError as e:
        st.error(f"❌ {e}"); st.stop()

    for w in warnings:
        st.markdown(f'<div class="warn-box">⚠️ {w}</div>', unsafe_allow_html=True)

    st.markdown(f"**{len(persons)} people loaded.** Preview:")
    st.dataframe(df_orig.head(5), use_container_width=True, hide_index=True)

    # ── Keyword scanner UI ───────────────────────────────────────────────────
    st.markdown("### 2. Set keywords to scan (optional)")
    st.markdown('''<div class="info-box">
Enter words or phrases to flag — e.g. <b>MAGA, Trump, January 6, Stop the Steal</b>.
The app will check each matched profile's <b>bio AND recent tweets</b> for these terms.
Leave blank to skip keyword scanning.
</div>''', unsafe_allow_html=True)

    kw_input = st.text_input(
        "Keywords (comma-separated)",
        placeholder="e.g. MAGA, Trump, January 6, Stop the Steal",
        label_visibility="collapsed",
    )
    keywords = [k.strip() for k in kw_input.split(",") if k.strip()] if kw_input else []

    if keywords:
        pills = " ".join(f'<span class="breakdown-pill" style="color:#fbbf24;background:#71350030;border:1px solid #d9770040">{k}</span>' for k in keywords)
        st.markdown(f'<div style="margin:0.4rem 0">{pills}</div>', unsafe_allow_html=True)
        st.markdown(f'<div class="info-box">⚠️ Keyword scanning adds ~1 extra API call per matched profile. For 100 profiles at ~80% match rate, expect ~80 extra calls (~$1.20 extra cost).</div>', unsafe_allow_html=True)

    st.markdown("### 3. Run the search")
    no_key = not anthropic_key and not serpapi_key
    if no_key:
        st.markdown('<div class="warn-box">⚠️ Add your Anthropic API key in the sidebar to enable search.</div>', unsafe_allow_html=True)

    col1, col2 = st.columns([3,1])
    with col1:
        run = st.button("🔍 Find Twitter Profiles", type="primary",
                         use_container_width=True, disabled=no_key)
    with col2:
        dry = st.button("🧪 Test (first 3)", use_container_width=True, disabled=no_key)

    if run or dry:
        sample  = persons[:3] if dry else persons
        results = []
        progress   = st.progress(0, text="Starting...")
        status_box = st.empty()
        live       = st.empty()

        for i, person in enumerate(sample):
            status_box.markdown(
                f'<div class="info-box">Searching for <b>{person.full_name}</b> ({i+1} of {len(sample)})...</div>',
                unsafe_allow_html=True)
            result = find_twitter(person, anthropic_key=anthropic_key, serpapi_key=serpapi_key)

            # Keyword scan — only for profiles we actually found
            if keywords and result.profile and result.confidence != "NO MATCH":
                status_box.markdown(
                    f'<div class="info-box">Scanning tweets for <b>{person.full_name}</b> ({result.profile.handle})...</div>',
                    unsafe_allow_html=True)
                result = scan_keywords(result, keywords, anthropic_key)

            results.append(result)
            progress.progress((i+1)/len(sample), text=f"{i+1}/{len(sample)} complete")

            rows_html = ""
            for r in results:
                p = r.profile
                conf_cls   = r.confidence.replace(" ","")
                bar_color  = {"HIGH":"#4ade80","MEDIUM":"#fbbf24","LOW":"#fb923c","NO MATCH":"#f87171"}.get(r.confidence,"#6b7280")
                handle_html = (f'<a class="handle-link" href="{p.profile_url}" target="_blank">{p.handle}</a>'
                               if p else '<span style="color:#334155">—</span>')
                bd_html = " ".join(f'<span class="breakdown-pill">{k.replace("_"," ")} +{v}</span>'
                                   for k, v in r.score_breakdown.items())
                kw_flag_html = ""
                if r.keyword_flagged:
                    kw_pills = " ".join(
                        f'<span class="breakdown-pill" style="color:#fbbf24;background:#71350030;border:1px solid #d9770040">⚑ {kw}</span>'
                        for kw in r.keyword_hits.keys()
                    )
                    kw_flag_html = f'<div style="padding:0 1rem 0.4rem;font-size:0.72rem">🚩 <b style="color:#fbbf24">Keywords matched:</b> {kw_pills}</div>'
                elif keywords and r.profile and r.confidence != "NO MATCH":
                    kw_flag_html = '<div style="padding:0 1rem 0.4rem;font-size:0.72rem;color:#334155">✓ No keywords found</div>'

                rows_html += f"""
<div class="result-row">
  <span class="badge badge-{conf_cls}">{r.emoji} {r.confidence}</span>
  <span class="person-name">{r.person.full_name}</span>
  <span>{handle_html}</span>
  <div style="flex:1"><div class="score-bar-wrap"><div class="score-bar" style="width:{r.score}%;background:{bar_color}"></div></div></div>
  <span class="score-num">{r.score}</span>
</div>
{f'<div style="padding:0 1rem 0.25rem;font-size:0.72rem;color:#475569">{bd_html}</div>' if bd_html else ''}
{kw_flag_html}"""
            live.markdown(rows_html, unsafe_allow_html=True)
            time.sleep(0.1)

        status_box.empty(); progress.empty()

        counts = {"HIGH":0,"MEDIUM":0,"LOW":0,"NO MATCH":0}
        for r in results: counts[r.confidence] += 1
        kw_flagged_count = sum(1 for r in results if r.keyword_flagged)

        kw_metric = f'''<div class="metric"><div class="num" style="color:#fbbf24">{kw_flagged_count}</div><div class="lbl">Keyword hits</div></div>''' if keywords else ""

        st.markdown(f"""
<div class="metric-row">
  <div class="metric"><div class="num" style="color:#e2e8f0">{len(results)}</div><div class="lbl">Searched</div></div>
  <div class="metric"><div class="num" style="color:#4ade80">{counts['HIGH']}</div><div class="lbl">High</div></div>
  <div class="metric"><div class="num" style="color:#fbbf24">{counts['MEDIUM']}</div><div class="lbl">Medium</div></div>
  <div class="metric"><div class="num" style="color:#fb923c">{counts['LOW']}</div><div class="lbl">Low</div></div>
  <div class="metric"><div class="num" style="color:#f87171">{counts['NO MATCH']}</div><div class="lbl">No match</div></div>
  {kw_metric}
</div>""", unsafe_allow_html=True)

        st.markdown("### 4. Download results")
        uploaded.seek(0)
        out_buf  = write_results(uploaded, results)
        out_name = uploaded.name.replace(".csv","").replace(".xlsx","") + "_twitter_results.xlsx"
        st.download_button("⬇️ Download Annotated Spreadsheet", data=out_buf,
                           file_name=out_name,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           use_container_width=True, type="primary")
        if dry and len(persons) > 3:
            st.info(f"Test run complete — click 'Find Twitter Profiles' to run all {len(persons)} people.")

else:
    st.markdown("""
<div style="border:1px solid #1e293b;border-radius:10px;background:#0f172a;padding:3rem 2rem;text-align:center;margin-top:1rem">
  <div style="font-size:2.5rem;margin-bottom:1rem">📋</div>
  <div style="font-size:1rem;color:#475569;margin-bottom:0.5rem">No file uploaded yet</div>
  <div style="font-size:0.82rem;color:#334155">Upload an .xlsx or .csv with First Name, Last Name, Address, and Employer columns.</div>
</div>""", unsafe_allow_html=True)

    sample_df = pd.DataFrame({
        "First Name": ["Jane","Michael","Sarah"],
        "Last Name":  ["Smith","Chen","Johnson"],
        "Address":    ["123 Main St, New York, NY 10001","456 Oak Ave, San Francisco, CA 94102","789 Pine Rd, Chicago, IL 60601"],
        "Employer":   ["Goldman Sachs","Google","United Airlines"],
    })
    buf = io.BytesIO()
    sample_df.to_excel(buf, index=False)
    buf.seek(0)
    st.download_button("📥 Download Sample Template", data=buf,
                       file_name="twitter_finder_template.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
